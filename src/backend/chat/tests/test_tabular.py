"""Tests for chat.tabular: spreadsheet loading, type inference and structured queries."""

from datetime import datetime
from io import BytesIO

import openpyxl
import pytest

from chat.constants import XLSX_MIME_TYPE
from chat.tabular import (
    TableQuery,
    TabularError,
    TabularFile,
    build_dataset,
    describe_tables,
    run_table_query,
)

FRENCH_CSV = (
    "Région;Montant;Date de vente;Code postal\n"
    "Île-de-France;1 234,56;03/02/2024;01000\n"
    "Bretagne;12,5;15/02/2024;35000\n"
    "Bretagne;-0,5;29/03/2024;29200\n"
).encode("cp1252")

BUDGETS_CSV = "region,budget\nBretagne,100\nÎle-de-France,1000\n".encode()


def _csv(content: bytes, file_name: str = "ventes.csv") -> TabularFile:
    return TabularFile(file_name=file_name, content_type="text/csv", content=content)


def _xlsx(sheets: dict[str, list[list]], file_name: str = "budget.xlsx") -> TabularFile:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return TabularFile(file_name=file_name, content_type=XLSX_MIME_TYPE, content=buffer.getvalue())


def _run(dataset, max_rows=100, **query) -> str:
    return run_table_query(dataset, TableQuery(**query), max_rows=max_rows)


def test_french_csv_types_are_inferred():
    """`;` separator, cp1252, decimal commas, thousands spaces and dd/mm/yyyy dates load typed."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    [table] = dataset.tables
    assert table.name == "ventes"
    assert table.row_count == 3
    assert [(c.name, c.header, c.kind) for c in table.columns] == [
        ("region", "Région", "text"),
        ("montant", "Montant", "real"),
        ("date_de_vente", "Date de vente", "date"),
        ("code_postal", "Code postal", "text"),
    ]
    frame = dataset.frames["ventes"]
    assert frame["montant"].sum() == pytest.approx(1246.56)
    assert frame["date_de_vente"].max() == datetime(2024, 3, 29)
    assert frame["code_postal"].tolist() == ["01000", "35000", "29200"]


def test_comma_csv_reads_dot_decimals():
    """With a `,` separator, the dot is the decimal separator."""
    dataset = build_dataset([_csv(b"name,amount\nalpha,1.5\nbeta,10.25\n", "data.csv")])

    assert [c.kind for c in dataset.tables[0].columns] == ["text", "real"]
    assert dataset.frames["data"]["amount"].sum() == pytest.approx(11.75)


def test_mixed_column_stays_text():
    """A column with one non-numeric value is kept as text, values unchanged."""
    dataset = build_dataset([_csv(b"ref;qty\nA;1\nB;n/a\n", "stock.csv")])

    assert [c.kind for c in dataset.tables[0].columns] == ["text", "text"]
    assert dataset.frames["stock"]["qty"].tolist() == ["1", "n/a"]


def test_invalid_date_keeps_column_as_text():
    """31/02 is not a date: the column stays text instead of being mangled."""
    dataset = build_dataset([_csv(b"d\n01/02/2024\n31/02/2024\n", "dates.csv")])

    assert dataset.tables[0].columns[0].kind == "text"


def test_headers_are_sanitized_and_deduplicated():
    """Accented, duplicate, empty and digit-leading headers become valid identifiers."""
    dataset = build_dataset([_csv(b"Nom;Nom;;2024\na;b;c;d\n", "t.csv")])

    assert [c.name for c in dataset.tables[0].columns] == ["nom", "nom_2", "column_3", "c_2024"]


def test_rows_wider_than_header_get_extra_columns():
    """Cells beyond the header are kept under generated column names."""
    dataset = build_dataset([_csv(b"a;b\n1;2;3\n", "t.csv")])

    assert list(dataset.frames["t"].columns) == ["a", "b", "column_3"]


def test_integer_column_with_empty_cell_stays_integer():
    """An empty cell does not turn an integer column into floats."""
    dataset = build_dataset([_csv(b"n\n1\n\n3\n", "t.csv")])

    assert dataset.tables[0].columns[0].kind == "integer"
    assert _run(dataset, table="t", aggregations=[{"func": "sum", "column": "n"}]) == (
        "| n_sum |\n|---|\n| 4 |"
    )


def test_xlsx_sheets_become_tables():
    """Each sheet of a workbook is a table; native numbers and dates are kept."""
    dataset = build_dataset(
        [
            _xlsx(
                {
                    "Recettes": [["Mois", "Total"], [datetime(2024, 1, 1), 1200.5]],
                    "Dépenses": [["Poste", "Total"], ["Loyer", 800]],
                }
            )
        ]
    )

    assert [(t.name, t.source) for t in dataset.tables] == [
        ("budget_recettes", "budget.xlsx (sheet Recettes)"),
        ("budget_depenses", "budget.xlsx (sheet Dépenses)"),
    ]
    assert _run(dataset, table="budget_recettes") == (
        "| mois | total |\n|---|---|\n| 2024-01-01 | 1200.5 |"
    )


def test_xlsx_whole_numbers_stay_integers():
    """Excel stores every number as a float: whole values load as an integer column."""
    dataset = build_dataset([_xlsx({"Feuil1": [["n", "x"], [800, 1.5], [2, 3]]})])

    assert [c.kind for c in dataset.tables[0].columns] == ["integer", "real"]
    assert _run(dataset, table="budget", aggregations=[{"func": "sum", "column": "n"}]) == (
        "| n_sum |\n|---|\n| 802 |"
    )


def test_xlsx_chart_sheets_are_skipped_and_hidden_sheets_kept():
    """Chart sheets hold no rows; hidden sheets often hold reference data, so they stay."""
    workbook = openpyxl.Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["a"])
    workbook.active.append([1])
    workbook.create_chartsheet("Chart")
    hidden = workbook.create_sheet("Ref")
    hidden.append(["b"])
    hidden.append([2])
    hidden.sheet_state = "hidden"
    buffer = BytesIO()
    workbook.save(buffer)

    dataset = build_dataset(
        [TabularFile(file_name="book.xlsx", content_type=XLSX_MIME_TYPE, content=buffer.getvalue())]
    )

    assert [t.name for t in dataset.tables] == ["book_data", "book_ref"]
    assert not dataset.skipped


def test_unreadable_file_is_skipped():
    """A corrupt workbook is reported as skipped; other files stay queryable."""
    broken = TabularFile(file_name="broken.xlsx", content_type=XLSX_MIME_TYPE, content=b"nope")
    dataset = build_dataset([broken, _csv(b"a\n1\n", "ok.csv")])

    assert [t.name for t in dataset.tables] == ["ok"]
    assert dataset.skipped == ["broken.xlsx"]
    assert "These files could not be read: broken.xlsx" in describe_tables(dataset)


def test_same_file_names_get_unique_tables():
    """Two files with the same name map to distinct tables."""
    dataset = build_dataset([_csv(b"a\n1\n", "x.csv"), _csv(b"a\n2\n", "x.csv")])

    assert [t.name for t in dataset.tables] == ["x", "x_2"]


def test_count_without_group_by():
    """No group_by, no aggregation column: count the matching rows."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    result = _run(
        dataset,
        table="ventes",
        filters=[{"column": "region", "op": "==", "value": "Bretagne"}],
        aggregations=[{"func": "count"}],
    )

    assert result == "| count |\n|---|\n| 2 |"


def test_group_by_with_sum_and_sort():
    """Aggregates per group, sorted on the aggregate name."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    result = _run(
        dataset,
        table="ventes",
        group_by=[{"column": "region"}],
        aggregations=[{"func": "sum", "column": "montant"}, {"func": "count"}],
        sort=[{"column": "montant_sum", "descending": True}],
    )

    assert result == (
        "| region | montant_sum | count |\n|---|---|---|\n"
        "| Île-de-France | 1234.56 | 1 |\n"
        "| Bretagne | 12 | 2 |"
    )


def test_group_by_month_on_date_column():
    """A date column can be grouped by month."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    result = _run(
        dataset, table="ventes", group_by=[{"column": "date_de_vente", "period": "month"}]
    )

    assert result == "| date_de_vente_month | count |\n|---|---|\n| 2024-02 | 2 |\n| 2024-03 | 1 |"


def test_date_filter_with_iso_value():
    """Date filters take ISO text values."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    result = _run(
        dataset,
        table="ventes",
        filters=[{"column": "date_de_vente", "op": "<", "value": "2024-03-01"}],
        aggregations=[{"func": "count"}],
    )

    assert result == "| count |\n|---|\n| 2 |"


def test_leading_zero_code_filter():
    """Codes with a leading zero are matched as text."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    result = _run(
        dataset,
        table="ventes",
        filters=[{"column": "code_postal", "op": "in", "value": ["01000", "29200"]}],
        columns=["code_postal"],
    )

    assert result == "| code_postal |\n|---|\n| 01000 |\n| 29200 |"


def test_contains_is_case_insensitive():
    """`contains` matches a substring regardless of case, never as a regex."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    result = _run(
        dataset,
        table="ventes",
        filters=[{"column": "region", "op": "contains", "value": "BRET"}],
        aggregations=[{"func": "count"}],
    )

    assert result == "| count |\n|---|\n| 2 |"


def test_join_adds_the_other_table_columns():
    """An inner join exposes the second table's columns to grouping and aggregation."""
    dataset = build_dataset([_csv(FRENCH_CSV), _csv(BUDGETS_CSV, "budgets.csv")])

    result = _run(
        dataset,
        table="ventes",
        join={"table": "budgets", "left_on": "region", "right_on": "region"},
        group_by=[{"column": "region"}],
        aggregations=[{"func": "sum", "column": "montant"}, {"func": "max", "column": "budget"}],
    )

    assert result == (
        "| region | montant_sum | budget_max |\n|---|---|---|\n"
        "| Bretagne | 12 | 100 |\n"
        "| Île-de-France | 1234.56 | 1000 |"
    )


def test_rows_are_capped():
    """Results beyond the limit are cut, with a note telling the model to aggregate."""
    dataset = build_dataset([_csv(b"n\n" + b"\n".join(str(i).encode() for i in range(10)))])

    result = _run(dataset, max_rows=3, table="ventes")

    assert result.count("\n| ") == 3
    assert "3 of 10 rows shown. The list is incomplete" in result


def test_limit_chosen_by_the_model_is_not_reported_as_incomplete():
    """A top-N limit set on purpose only states the total, without the 'incomplete' warning."""
    dataset = build_dataset([_csv(b"n\n" + b"\n".join(str(i).encode() for i in range(10)))])

    result = _run(dataset, table="ventes", sort=[{"column": "n", "descending": True}], limit=2)

    assert result == "| n |\n|---|\n| 9 |\n| 8 |\n\n2 of 10 rows shown."


@pytest.mark.parametrize(
    "query,message",
    [
        ({"table": "missing"}, 'Unknown table "missing". Available tables: ventes.'),
        (
            {"table": "ventes", "columns": ["montnat"]},
            "Unknown column(s) montnat in columns. Available columns: region, montant,",
        ),
        (
            {"table": "ventes", "aggregations": [{"func": "sum", "column": "region"}]},
            '"region" is not numeric: sum is not possible.',
        ),
        (
            {"table": "ventes", "group_by": [{"column": "region", "period": "month"}]},
            '"region" is not a date column',
        ),
        (
            {"table": "ventes", "filters": [{"column": "region", "op": "=="}]},
            'Filter on "region" with op == needs a value.',
        ),
        (
            {"table": "ventes", "aggregations": [{"func": "sum"}]},
            "Aggregation sum needs a column.",
        ),
    ],
)
def test_invalid_queries_explain_the_fix(query, message):
    """Errors name what is wrong and what is available, so the model can retry."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    with pytest.raises(TabularError) as exc_info:
        run_table_query(dataset, TableQuery(**query), max_rows=100)

    assert message in str(exc_info.value)


def test_type_mismatch_in_filter_is_reported():
    """Comparing a text column with a number is a model-fixable error, not a crash."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    with pytest.raises(TabularError, match="Invalid query"):
        _run(dataset, table="ventes", filters=[{"column": "region", "op": ">", "value": 3}])


def test_describe_tables_lists_columns_and_samples():
    """The description gives table names, headers as written, types and sample rows."""
    dataset = build_dataset([_csv(FRENCH_CSV)])

    description = describe_tables(dataset)

    assert "## Table ventes" in description
    assert "Rows: 3" in description
    assert '- montant (real) - header in file: "Montant"' in description
    assert "| Île-de-France | 1234.56 | 2024-02-03 | 01000 |" in description
