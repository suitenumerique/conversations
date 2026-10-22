"""
Tests for building a deck out of the user's own slides.

Real components: python-pptx. The user's deck is made with python-pptx's own
default template, built in memory and read back as a user would open it.
"""

from io import BytesIO
from pathlib import Path

import pptx
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from chat.file_generation import (
    PresentationStyle,
    TemplatedPresentation,
    TemplatedSlide,
    build_from_template,
    outline_template,
)

PICTURE = Path(__file__).parents[3] / "core/static/images/logo.png"


def user_deck() -> bytes:
    """A cover, then a content slide holding a picture and hand-sized text."""
    deck = pptx.Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = "Titre de la présentation"
    cover.placeholders[1].text = "Direction - date"

    content = deck.slides.add_slide(deck.slide_layouts[1])
    content.shapes.title.text = "Titre de contenu"
    content.placeholders[1].text = "Texte d'exemple"
    content.placeholders[1].text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    content.shapes.add_picture(str(PICTURE), Inches(8), Inches(6), Inches(1))

    blob = BytesIO()
    deck.save(blob)
    return blob.getvalue()


def build(*slides, style=None):
    """Build a deck out of the user's slides and reopen it."""
    presentation = TemplatedPresentation(title="Deck", slides=list(slides))
    blob = build_from_template(user_deck(), presentation, style or PresentationStyle())
    return pptx.Presentation(BytesIO(blob))


def texts(slide):
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def test_the_outline_describes_each_slide():
    """The agent sees each slide's text shapes, their ids and texts, and other content."""
    outline = outline_template(user_deck())

    assert outline.shape_ids == (frozenset({2, 3}), frozenset({2, 3}))
    assert "Slide 1, layout 'Title Slide':" in outline.description
    assert "- shape 2 (center_title): 'Titre de la présentation'" in outline.description
    assert '- shape 3 (object): "Texte d\'exemple"' in outline.description
    assert "- also holds: picture" in outline.description


def test_slides_are_copied_with_new_texts_and_the_originals_dropped():
    """Each new slide copies a user's slide, keeps its picture, and gets its new texts."""
    deck = build(
        TemplatedSlide(template_slide=1, texts={2: "Plan vélo", 3: "Comité"}),
        TemplatedSlide(template_slide=2, texts={2: "Constats", 3: "- un\n- deux"}),
        TemplatedSlide(template_slide=2, texts={2: "Suite", 3: "Trois"}),
    )

    assert [texts(slide) for slide in deck.slides] == [
        ["Plan vélo", "Comité"],
        ["Constats", "un\ndeux"],
        ["Suite", "Trois"],
    ]
    for slide in list(deck.slides)[1:]:
        picture = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        assert picture.image.blob == PICTURE.read_bytes()


def test_an_empty_text_clears_a_shape():
    deck = build(TemplatedSlide(template_slide=2, texts={2: "Titre de contenu", 3: ""}))

    assert texts(deck.slides[0]) == ["Titre de contenu", ""]


def test_the_look_of_hand_formatted_text_is_kept_under_the_style():
    """New text keeps the size set by hand on the text it replaces, and takes the style."""
    deck = build(
        TemplatedSlide(template_slide=2, texts={2: "Constats", 3: "Nouveau texte"}),
        style=PresentationStyle(text_color="#000091"),
    )

    run = deck.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    assert run.font.size == Pt(20)
    assert str(run.font.color.rgb) == "000091"


def test_presenter_notes_are_attached():
    deck = build(TemplatedSlide(template_slide=1, slide_notes="À dire à l'oral"))

    assert deck.slides[0].notes_slide.notes_text_frame.text == "À dire à l'oral"


def test_the_outline_rejects_slides_and_shapes_the_user_deck_does_not_have():
    """The agent is told what it copied that does not exist, to correct it."""
    outline = outline_template(user_deck())

    def check(*slides):
        return outline.check(TemplatedPresentation(title="Deck", slides=list(slides)))

    assert check(TemplatedSlide(template_slide=2, texts={2: "ok", 3: "ok"})) is None
    assert "copies slide 5" in check(TemplatedSlide(template_slide=5))
    assert "shapes [9]" in check(TemplatedSlide(template_slide=1, texts={2: "", 3: "", 9: "?"}))
    assert "no text to shapes [3]" in check(TemplatedSlide(template_slide=1, texts={2: "ok"}))


def deck_with_table_group_and_chart() -> bytes:
    """One slide holding a table, a group of two labelled tiles, and a chart."""
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(0.5), Inches(4), Inches(1)).table
    for row, values in enumerate([("Poste", "Montant"), ("Fonctionnement", "820 k€")]):
        table.cell(row, 0).text, table.cell(row, 1).text = values
    group = slide.shapes.add_group_shape()
    for index, label in enumerate(["Projet A", "Projet B"]):
        tile = group.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.5 + 2 * index), Inches(2), Inches(1.5), Inches(1)
        )
        tile.text_frame.text = label
    chart_data = CategoryChartData()
    chart_data.categories = ["2025"]
    chart_data.add_series("Visites", (4210,))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(0.5), Inches(4), Inches(3), chart_data
    )
    blob = BytesIO()
    deck.save(blob)
    return blob.getvalue()


def test_tables_and_grouped_texts_are_described_and_filled():
    """A table is shown and filled as a Markdown table; texts inside groups change too."""
    blob = deck_with_table_group_and_chart()
    outline = outline_template(blob)
    table_id, *tile_ids = sorted(outline.shape_ids[0])

    assert "(table): '| Poste | Montant |\\n|---|---|\\n| Fonctionnement | 820 k€ |'" in (
        outline.description
    )
    assert "'Projet A'" in outline.description

    presentation = TemplatedPresentation(
        title="Deck",
        slides=[
            TemplatedSlide(
                template_slide=1,
                texts={
                    table_id: "| Poste | Montant |\n|---|---|\n| Fonctionnement | 850 k€ |",
                    tile_ids[0]: "Portail v2",
                    tile_ids[1]: "Archives orales",
                },
            )
        ],
    )
    slide = pptx.Presentation(
        BytesIO(build_from_template(blob, presentation, PresentationStyle()))
    ).slides[0]

    table = next(shape for shape in slide.shapes if shape.has_table).table
    assert table.cell(1, 1).text == "850 k€"
    group = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
    assert [tile.text_frame.text for tile in group.shapes] == ["Portail v2", "Archives orales"]


def test_slides_copying_a_chart_are_reported():
    """A copied chart keeps its data: the slides that copy one are listed."""
    outline = outline_template(deck_with_table_group_and_chart())
    presentation = TemplatedPresentation(
        title="Deck", slides=[TemplatedSlide(template_slide=1), TemplatedSlide(template_slide=1)]
    )

    assert "- also holds: chart" in outline.description
    assert outline.copied_charts(presentation) == [1, 2]
