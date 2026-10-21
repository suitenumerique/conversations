"""
Tests for the deck-wide style.

Real components: the bundled pptx template and python-pptx. Decks are built in
memory and read back to assert what a user would actually open.
"""

from io import BytesIO

import pptx
import pytest
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt
from pydantic import ValidationError

from chat.file_generation import (
    GENERIC_TEMPLATE,
    Presentation,
    PresentationStyle,
    Slide,
    SlideType,
    build_presentation,
)

STYLE = PresentationStyle(
    title_color="#000091",
    text_color="E1000F",
    background_color="#F5F5FE",
    title_font="Arial",
    text_font="Georgia",
    title_size=32,
    text_size=18,
)


def build_deck(slides, style=None):
    """Build a deck from slides and reopen it as python-pptx would."""
    blob = build_presentation(GENERIC_TEMPLATE, Presentation(title="Deck", slides=slides), style)
    return pptx.Presentation(BytesIO(blob))


def runs_of(shape):
    """Return the runs of a text shape, or of every cell of a table."""
    frames = (
        [shape.text_frame]
        if shape.has_text_frame
        else [cell.text_frame for cell in shape.table.iter_cells()]
    )
    return [run for frame in frames for paragraph in frame.paragraphs for run in paragraph.runs]


def content_slide(**fields):
    """A one-column slide, after the cover every deck opens with."""
    return [
        Slide(type=SlideType.COVER, title="Couverture"),
        Slide(type=SlideType.TITLE_ONE_COLUMN, title="Constats", **fields),
    ]


def font_scale(shape):
    """The shrink-to-fit scale stored on a shape, in percent, if any."""
    autofit = shape.text_frame._bodyPr.find(qn("a:normAutofit"))  # pylint: disable=protected-access
    scale = autofit.get("fontScale") if autofit is not None else None
    return int(scale) / 1000 if scale else None


def test_titles_and_text_are_styled_apart():
    """Titles take the title style, every other field the text style."""
    deck = build_deck(content_slide(content="- Premier\n- **Second**"), STYLE)
    title_shape, body_shape = deck.slides[1].shapes[0], deck.slides[1].shapes[1]

    for run in runs_of(title_shape):
        assert run.font.color.rgb == RGBColor.from_string("000091")
        assert run.font.name == "Arial"
        assert run.font.size == Pt(32)
    for run in runs_of(body_shape):
        assert run.font.color.rgb == RGBColor.from_string("E1000F")
        assert run.font.name == "Georgia"
        assert run.font.size == Pt(18)


def test_cover_and_section_slides_keep_their_layout_sizes():
    """
    Cover and section slides, designed with larger text, take the colours and
    fonts but not the sizes.
    """
    deck = build_deck(
        [
            Slide(type=SlideType.COVER, title="Couverture", subtitle="Sous-titre"),
            Slide(type=SlideType.SECTION, title="Partie", subtitle="Détail"),
        ],
        STYLE,
    )

    runs = [run for slide in deck.slides for shape in slide.shapes for run in runs_of(shape)]
    assert {run.text for run in runs} == {"Couverture", "Sous-titre", "Partie", "Détail"}
    for run in runs:
        assert run.font.size is None
        assert run.font.name in ("Arial", "Georgia")
        assert run.font.color.rgb in (
            RGBColor.from_string("000091"),
            RGBColor.from_string("E1000F"),
        )


def test_tables_take_the_text_font_but_keep_their_colours():
    """A table takes the text font, but keeps the colours designed for its cells' fills."""
    deck = build_deck(content_slide(content="| A | B |\n|---|---|\n| 1 | 2 |"), STYLE)
    table_shape = next(shape for shape in deck.slides[1].shapes if shape.has_table)

    runs = runs_of(table_shape)
    assert runs
    assert all(run.font.color.type is None for run in runs)
    assert all(run.font.name == "Georgia" for run in runs)


def test_a_long_table_gets_a_smaller_text_size():
    """A table never shrinks on its own: a long one is given a smaller size to fit."""
    short = "| A | B |\n|---|---|\n| 1 | 2 |"
    long = "| A | B |\n|---|---|\n" + "\n".join(
        f"| Indicateur numéro {i} au libellé assez long | Valeur {i} |" for i in range(12)
    )

    def table_size(content):
        deck = build_deck(content_slide(content=content), PresentationStyle(text_size=16))
        table_shape = next(shape for shape in deck.slides[1].shapes if shape.has_table)
        return {run.font.size for run in runs_of(table_shape)}

    assert table_size(short) == {Pt(16)}
    (long_size,) = table_size(long)
    assert Pt(8) <= long_size < Pt(16)


def test_inline_code_keeps_its_own_font():
    """A font the renderer set on purpose is not overridden by the style."""
    deck = build_deck(content_slide(content="Lancer `make run`"), STYLE)

    code_run = next(run for run in runs_of(deck.slides[1].shapes[1]) if run.text == "make run")
    assert code_run.font.name == "Consolas"
    assert code_run.font.size == Pt(10)


def test_the_background_is_filled():
    """Every slide gets the background colour."""
    deck = build_deck(content_slide(content="Corps"), STYLE)

    for slide in deck.slides:
        assert slide.background.fill.fore_color.rgb == RGBColor.from_string("F5F5FE")


def test_without_style_the_template_styling_is_kept():
    """No style sets nothing, leaving every run to the template."""
    deck = build_deck(content_slide(content="Corps"))

    for shape in deck.slides[1].shapes:
        if shape.has_text_frame:
            for run in runs_of(shape):
                assert run.font.size is None
                assert run.font.name is None
    assert deck.slides[1].follow_master_background


def test_a_partial_style_only_sets_its_fields():
    """Fields left unset keep the template's value."""
    deck = build_deck(content_slide(content="Corps"), PresentationStyle(text_size=20))
    title_shape, body_shape = deck.slides[1].shapes[0], deck.slides[1].shapes[1]

    assert all(run.font.size is None for run in runs_of(title_shape))
    assert all(run.font.size == Pt(20) for run in runs_of(body_shape))
    assert all(run.font.name is None for run in runs_of(body_shape))


@pytest.mark.parametrize(
    "field,value",
    [
        ("title_color", "blue"),
        ("text_color", "#12345"),
        ("title_size", 200),
        ("text_size", 2),
        ("text_font", "x" * 65),
    ],
)
def test_invalid_style_values_are_rejected(field, value):
    """Colours must be hex codes, sizes and font names stay in sane bounds."""
    with pytest.raises(ValidationError):
        PresentationStyle(**{field: value})


def test_updated_with_only_overrides_the_fields_set():
    """Updating a style keeps every field the changes leave unset."""
    current = PresentationStyle(title_color="#000091", text_size=14)

    updated = current.updated_with(PresentationStyle(text_size=20, text_font="Arial"))

    assert updated == PresentationStyle(title_color="#000091", text_size=20, text_font="Arial")


def test_text_too_long_for_its_shape_is_scaled_down():
    """Overflowing text stores a shrink scale, which PowerPoint applies on opening."""
    long_text = "\n".join(
        f"- Une puce assez longue pour occuper toute une ligne, numéro {i}" for i in range(15)
    )
    deck = build_deck(content_slide(content=long_text), PresentationStyle(text_size=20))

    scale = font_scale(deck.slides[1].shapes[1])
    assert scale is not None
    assert 50 <= scale < 100


def test_text_that_fits_is_not_scaled():
    """Text that fits its shape keeps its size."""
    deck = build_deck(content_slide(content="- Court\n- Bref"), PresentationStyle(text_size=20))

    assert font_scale(deck.slides[1].shapes[1]) is None
