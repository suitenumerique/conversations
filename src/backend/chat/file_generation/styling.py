"""Apply a deck-wide style to the slides once their content is rendered.

The style is set on the runs the renderer wrote, rather than on the template's
theme or masters: the template spreads sizes and colours across its master and
each layout, so overriding the written text is the one place that reliably wins.
"""

import math

from pptx.dml.color import RGBColor
from pptx.shapes.base import BaseShape
from pptx.slide import Slide as PptxSlide
from pptx.util import Emu, Pt

from chat.file_generation.entities import PresentationStyle
from chat.file_generation.renderer import INDENT_PER_LEVEL

# The slide field styled as a title; every other field is styled as text.
TITLE_FIELD = "title"

# Rough text metrics, to estimate whether text fits its shape: the template's
# body text size, and glyph width, line height and paragraph gap in font sizes.
DEFAULT_TEXT_SIZE = 14
CHAR_WIDTH = 0.55
LINE_HEIGHT = 1.2
PARAGRAPH_GAP = 0.5
# Default inner margins of a text shape, and the smallest scale text shrinks to.
HORIZONTAL_INSET = Emu(91440).pt
VERTICAL_INSET = Emu(45720).pt
SMALLEST_SCALE = 0.5


def apply_background(pptx_slide: PptxSlide, style: PresentationStyle) -> None:
    """Fill the slide background with the style's colour, if it sets one."""
    if style.background_color:
        fill = pptx_slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(style.background_color)


def apply_text_style(shape: BaseShape, field: str, style: PresentationStyle) -> None:
    """
    Style the text a shape holds, as a title or as body text depending on ``field``,
    and shrink body text so that it fits the shape.

    A font or size the renderer set on purpose (inline code) is kept, and a
    table keeps its colours, designed for its cells' fills.
    """
    if field == TITLE_FIELD:
        color, font, size = style.title_color, style.title_font, style.title_size
    else:
        color, font, size = style.text_color, style.text_font, style.text_size

    for run in _runs(shape):
        if color and not shape.has_table:
            run.font.color.rgb = _rgb(color)
        if font and run.font.name is None:
            run.font.name = font
        if size and run.font.size is None:
            run.font.size = Pt(size)

    if field != TITLE_FIELD:
        _shrink_to_fit(shape, size or DEFAULT_TEXT_SIZE)


def _shrink_to_fit(shape: BaseShape, size: float) -> None:
    """
    Shrink the text of a shape so that it fits it.

    PowerPoint only shrinks text to fit once it is edited: on opening, it uses
    the scale stored in the file, so that scale is estimated here. Tables never
    shrink on their own, so their text is given the reduced size directly.
    """
    if shape.has_table:
        table = shape.table
        rows = [
            [
                (cell.text, column.width.pt - 2 * HORIZONTAL_INSET)
                for cell, column in zip(row.cells, table.columns, strict=True)
            ]
            for row in table.rows
        ]
        font_size = Pt(size * _fitting_scale(rows, shape.height.pt, size))
        for run in _runs(shape):
            run.font.size = font_size
    else:
        width = shape.width.pt - 2 * HORIZONTAL_INSET
        rows = [
            [(paragraph.text, width - Emu(INDENT_PER_LEVEL).pt * (paragraph.level + 1))]
            for paragraph in shape.text_frame.paragraphs
        ]
        scale = _fitting_scale(rows, shape.height.pt - 2 * VERTICAL_INSET, size)
        if scale < 1:
            # pylint: disable=protected-access
            autofit = shape.text_frame._bodyPr.get_or_change_to_normAutofit()  # noqa: SLF001
            autofit.set("fontScale", str(round(scale * 100_000)))


def _fitting_scale(rows: list[list[tuple[str, float]]], height: float, size: float) -> float:
    """
    Estimate the scale, down to SMALLEST_SCALE, at which rows of text fit a height.

    Each row lays texts side by side, each in its own width, as table cells do;
    a text shape is a column of one-paragraph rows. Sizes are in points.
    """
    scale = 1.0
    while scale > SMALLEST_SCALE and sum(_row_height(row, size * scale) for row in rows) > height:
        scale -= 0.05
    return scale


def _row_height(row: list[tuple[str, float]], font_size: float) -> float:
    """Estimate the height of texts laid side by side, the tallest one setting it."""
    line_count = max(
        math.ceil(len(text) * CHAR_WIDTH * font_size / width) or 1 for text, width in row
    )
    return (line_count * LINE_HEIGHT + PARAGRAPH_GAP) * font_size


def _runs(shape: BaseShape):
    """Yield the runs of a text shape, or of every cell of a table."""
    if shape.has_table:
        text_frames = [cell.text_frame for cell in shape.table.iter_cells()]
    else:
        text_frames = [shape.text_frame]
    for text_frame in text_frames:
        for paragraph in text_frame.paragraphs:
            yield from paragraph.runs


def _rgb(hex_color: str) -> RGBColor:
    """Parse a ``#RRGGBB`` or ``RRGGBB`` colour."""
    return RGBColor.from_string(hex_color.removeprefix("#"))
