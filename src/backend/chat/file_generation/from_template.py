"""Build a deck out of the user's own slides, so that it keeps their design.

Each slide of the new deck is a copy of one of the user's slides, or a new
slide made from one of their layouts, whose text shapes and tables may get new
Markdown texts; everything else on the slide (master, layout, pictures, charts,
the shapes given no text) is kept as it is. The user's slides that are not
copied are dropped.
"""

import copy
import json
import re
from collections import Counter
from dataclasses import dataclass
from io import BytesIO

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.base import BaseShape

from chat.file_generation.entities import PresentationStyle, TemplatedPresentation
from chat.file_generation.renderer import MONO_FONT, render, render_into_table
from chat.file_generation.styling import TITLE_FIELD, apply_text_style

# Attributes in this namespace hold relationship ids, e.g. r:embed on a picture.
RELATIONSHIP_ATTRIBUTE = qn("r:id").removesuffix("id")
# What a copy must not bring along: links to other slides, which would keep the
# slides the deck drops inside its file, and review comments, legacy and modern.
NOT_COPIED = (
    RT.SLIDE,
    RT.COMMENTS,
    "http://schemas.microsoft.com/office/2018/10/relationships/comments",
)
# Paragraph properties that make a paragraph a bulleted or numbered list item.
LIST_MARKERS = (qn("a:buAutoNum"), qn("a:buChar"))
# The number the renderer writes as text at the start of a numbered list item.
RENDERED_NUMBER = re.compile(r"\d+\. ")


@dataclass(frozen=True)
class TemplateOutline:
    """What the presentation agent needs to know of the user's deck."""

    # The ids of the text shapes and tables of each slide, to check the agent's answer.
    shape_ids: tuple[frozenset[int], ...]
    # The ids of the placeholders of a new slide made from each layout, by layout name.
    layout_shape_ids: dict[str, frozenset[int]]
    # The slides and layouts described for the agent's prompt.
    description: str

    def check(self, presentation: TemplatedPresentation) -> str | None:
        """Explain how a deck does not fit the user's slides, if it does not."""
        copies = Counter(slide.template_slide for slide in presentation.slides)
        for number, slide in enumerate(presentation.slides, start=1):
            if slide.layout:
                if slide.layout not in self.layout_shape_ids:
                    return f"Slide {number} uses layout {slide.layout!r}, which does not exist."
                shape_ids, source = self.layout_shape_ids[slide.layout], f"layout {slide.layout!r}"
            elif slide.template_slide and 1 <= slide.template_slide <= len(self.shape_ids):
                shape_ids = self.shape_ids[slide.template_slide - 1]
                source = f"slide {slide.template_slide}"
            else:
                return f"Slide {number} copies slide {slide.template_slide}, which does not exist."
            if unknown := set(slide.texts) - shape_ids:
                return (
                    f"Slide {number} gives texts to shapes {sorted(unknown)}, which {source} "
                    f"does not have: its shapes are {sorted(shape_ids)}."
                )
            copied_again = not slide.layout and copies[slide.template_slide] > 1
            if copied_again and (missing := shape_ids - set(slide.texts)):
                return (
                    f"Slide {number} copies slide {slide.template_slide}, which is copied several "
                    f"times, but gives no text to shapes {sorted(missing)}: give each copy a text "
                    "for every shape, an empty one to remove it."
                )
        return None


def outline_template(blob: bytes) -> TemplateOutline:
    """Describe the user's slides: their text shapes, current texts and other content."""
    deck = pptx.Presentation(BytesIO(blob))
    shape_ids, lines = [], []
    for number, slide in enumerate(deck.slides, start=1):
        shapes = _text_shapes(slide.shapes)
        shape_ids.append(frozenset(shape.shape_id for shape in shapes))
        lines.append(f"Slide {number}, layout {slide.slide_layout.name!r}:")
        for shape in shapes:
            # As JSON, which keeps characters like non-breaking spaces as they are.
            text = json.dumps(_shape_text(shape), ensure_ascii=False)
            lines.append(f"- {_describe(shape)}: {text}")
        if pictures_and_charts := sorted(
            {_non_text_content(shape) for shape in slide.shapes} - {None}
        ):
            lines.append(f"- also holds: {', '.join(pictures_and_charts)}")

    layout_shape_ids = {}
    for name, layout in _layouts(deck).items():
        if shapes := _text_shapes(deck.slides.add_slide(layout).shapes):
            layout_shape_ids[name] = frozenset(shape.shape_id for shape in shapes)
            lines.append(f"Layout {name!r}:")
            lines += [f"- {_describe(shape)}" for shape in shapes]
    return TemplateOutline(
        shape_ids=tuple(shape_ids),
        layout_shape_ids=layout_shape_ids,
        description="\n".join(lines),
    )


def build_from_template(
    blob: bytes, presentation: TemplatedPresentation, style: PresentationStyle
) -> bytes:
    """Render a deck made of copies of the user's slides into the bytes of a pptx file."""
    deck = pptx.Presentation(BytesIO(blob))
    template_slides = list(deck.slides)
    template_slide_ids = list(deck.slides._sldIdLst)  # noqa: SLF001 # pylint: disable=protected-access

    for templated in presentation.slides:
        if templated.layout:
            slide = deck.slides.add_slide(_layouts(deck)[templated.layout])
        else:
            slide = _copy_slide(deck, template_slides[templated.template_slide - 1])
        for shape in _text_shapes(slide.shapes):
            text = templated.texts.get(shape.shape_id)
            if text is None or text.split() == _shape_text(shape).split():
                continue
            if not text:
                shape._element.getparent().remove(shape._element)  # noqa: SLF001 # pylint: disable=protected-access
            elif shape.has_table:
                render_into_table(shape.table, text)
            else:
                looks = _looks(shape)
                content = render(shape, text)
                _restore_looks(shape, looks)
                apply_text_style(content, TITLE_FIELD if _is_title(shape) else "text", style)
        if templated.slide_notes:
            render(slide.notes_slide.notes_text_frame, templated.slide_notes)

    for slide_id in template_slide_ids:
        deck.slides._sldIdLst.remove(slide_id)  # noqa: SLF001 # pylint: disable=protected-access
        deck.part.drop_rel(slide_id.rId)

    # The user's deck may be shared: do not hand over who wrote it.
    properties = deck.core_properties
    for name in ("author", "category", "comments", "keywords", "last_modified_by", "subject"):
        setattr(properties, name, "")
    properties.title = presentation.title[:255]

    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def _copy_slide(deck, source):
    """Append a copy of a slide, with its shapes, background and the parts they use."""
    slide = deck.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)  # noqa: SLF001 # pylint: disable=protected-access

    # Pictures and charts point to other parts through relationships: relate
    # the copy to the same parts, under the ids it gives them.
    rels = source.part.rels.values()
    new_ids = {
        rel.rId: slide.part.relate_to(
            rel.target_ref if rel.is_external else rel.target_part,
            rel.reltype,
            is_external=rel.is_external,
        )
        for rel in rels
        if rel.reltype not in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, *NOT_COPIED)
    }
    dropped_ids = {rel.rId for rel in rels if rel.reltype in NOT_COPIED}
    for source_element in source.shapes._spTree.iter_shape_elms():  # noqa: SLF001 # pylint: disable=protected-access
        element = copy.deepcopy(source_element)
        for node in list(element.iter()):
            for name, value in list(node.attrib.items()):
                if not name.startswith(RELATIONSHIP_ATTRIBUTE):
                    continue
                if value in dropped_ids:  # A link to another slide, gone with it.
                    node.getparent().remove(node)
                    break
                if value in new_ids:
                    node.set(name, new_ids[value])
        slide.shapes._spTree.append(element)  # noqa: SLF001 # pylint: disable=protected-access

    if not source.follow_master_background:
        slide._element.cSld.insert(0, copy.deepcopy(source._element.cSld.bg))  # noqa: SLF001 # pylint: disable=protected-access
    return slide


def _layouts(deck) -> dict:
    """The layouts of every master of a deck, by name, the first one winning."""
    layouts = {}
    for master in deck.slide_masters:
        for layout in master.slide_layouts:
            layouts.setdefault(layout.name, layout)
    return layouts


def _text_shapes(shapes) -> list[BaseShape]:
    """The shapes holding text or a table, inside groups too, in reading order."""
    found = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found += _text_shapes(shape.shapes)
        elif shape.has_text_frame or shape.has_table:
            found.append(shape)
    return found


def _shape_text(shape: BaseShape) -> str:
    """The text of a shape, a table being written as a Markdown table."""
    if not shape.has_table:
        return shape.text_frame.text
    rows = [
        "| " + " | ".join(cell.text.replace("\n", "<br>") for cell in row.cells) + " |"
        for row in shape.table.rows
    ]
    separator = "|" + "---|" * len(shape.table.columns)
    return "\n".join([rows[0], separator, *rows[1:]])


def _looks(shape: BaseShape) -> dict:
    """
    The properties of the first paragraph and run of each level of a shape, or of
    its layout's sample title when it is an empty title, for its new text to take.
    """
    if shape.text_frame.text.strip():
        source = shape
    else:  # The sample text of a body is filler, often in bold: only a title's is a look.
        source = _layout_sample(shape) if _is_title(shape) else None
    looks = {}
    for paragraph in source.text_frame.paragraphs if source else []:
        if paragraph.level in looks or not paragraph.runs:
            continue
        all_properties = [run._r.get_or_add_rPr() for run in paragraph.runs]  # noqa: SLF001 # pylint: disable=protected-access
        run_properties = copy.deepcopy(all_properties[0])
        # Emphasis is the paragraph's look only when all its runs share it, not a bold label.
        for emphasis in ("b", "i", "u"):
            if len({properties.get(emphasis) for properties in all_properties}) > 1:
                run_properties.attrib.pop(emphasis, None)
        looks[paragraph.level] = (copy.deepcopy(paragraph._p.pPr), run_properties)  # noqa: SLF001 # pylint: disable=protected-access
    return looks


def _restore_looks(shape: BaseShape, looks: dict) -> None:
    """Give the new paragraphs and runs of a shape the properties it had at their level."""
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.level in looks:
            _restore_paragraph_look(paragraph, looks[paragraph.level][0])
        if look := looks.get(paragraph.level) or looks.get(0):
            for run in paragraph.runs:
                if run.font.name != MONO_FONT:
                    _merge_run_properties(run, look[1])


def _restore_paragraph_look(paragraph, properties) -> None:
    """
    Give a new paragraph the properties of the shape's paragraph at its level.

    Its bullets and numbering replace the renderer's, unless only the new text is a list.
    """
    element = paragraph._p  # noqa: SLF001 # pylint: disable=protected-access
    rendered, number = element.get_or_add_pPr(), _rendered_number(paragraph)
    if not _is_list(properties) and (_is_list(rendered) or number):
        return
    if number and properties.find(qn("a:buAutoNum")) is not None:
        element.remove(number._r)  # noqa: SLF001 # pylint: disable=protected-access
    element.remove(rendered)
    if properties is not None:
        element.insert(0, copy.deepcopy(properties))


def _merge_run_properties(run, look) -> None:
    """Set a run's properties to a copy of ``look``, under the ones the run sets itself."""
    own = run._r.get_or_add_rPr()  # noqa: SLF001 # pylint: disable=protected-access
    merged = copy.deepcopy(look)
    for name, value in own.attrib.items():
        merged.set(name, value)
    for child in own:  # Only a hyperlink: the renderer sets fonts on code runs alone.
        merged.append(child)
    own.getparent().replace(own, merged)


def _is_list(properties) -> bool:
    return properties is not None and any(properties.find(tag) is not None for tag in LIST_MARKERS)


def _rendered_number(paragraph):
    """The run holding the number the renderer wrote at the start of an item, if any."""
    runs = paragraph.runs
    return runs[0] if runs and RENDERED_NUMBER.fullmatch(runs[0].text) else None


def _layout_sample(shape: BaseShape):
    """
    The placeholder of a shape's layout, when the layout styles it on its sample
    text rather than in its own style, as LibreOffice does.
    """
    layout = shape._base_placeholder if shape.is_placeholder else None  # noqa: SLF001 # pylint: disable=protected-access
    if layout is None or len(layout._element.txBody.find(qn("a:lstStyle"))):  # noqa: SLF001 # pylint: disable=protected-access
        return None
    return layout


def _is_title(shape: BaseShape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in (
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
    )


def _describe(shape: BaseShape) -> str:
    """A shape as the agent sees it: its id, kind and size, for its text to fit it."""
    size = f"{shape.width.cm:.0f}×{shape.height.cm:.0f} cm"
    return f"shape {shape.shape_id} ({_shape_kind(shape)}, {size})"


def _shape_kind(shape: BaseShape) -> str:
    """How a shape is described to the agent: its placeholder type, a table or a text box."""
    if shape.has_table:
        return "table"
    if shape.is_placeholder:
        kind = shape.placeholder_format.type
        return "body" if kind == PP_PLACEHOLDER.OBJECT else kind.name.lower()
    return "text box"


def _non_text_content(shape: BaseShape) -> str | None:
    """The kind of content of a shape that copies keep as it is, if any."""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape.has_chart:
        return "chart"
    return None
